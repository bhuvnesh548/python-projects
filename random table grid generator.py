#random table grid generator 
from docx import Document
import random
doc = Document()
header_numbers = list(range(2, 13))
random.shuffle(header_numbers)
left_numbers = list(range(2, 20))
random.shuffle(left_numbers)
rows = len(left_numbers) + 1      
cols = len(header_numbers) + 1    
table = doc.add_table(rows=rows, cols=cols)
table.style = "Table Grid"              
table.cell(0, 0).text = "blank"
for j, num in enumerate(header_numbers, start=1):
    cell = table.cell(0, j)
    cell.text = str(num)
for i, num in enumerate(left_numbers, start=1):
    cell = table.cell(i, 0)
    cell.text = str(num)
doc.save("table.docx")
